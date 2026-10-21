# SPDX-FileCopyrightText: 2026 Sebastien Rousseau <sebastian.rousseau@gmail.com>
# SPDX-License-Identifier: Apache-2.0 OR MIT
"""Coverage-guided fuzzing of ``write_xlsx`` with Atheris.

The contract, the same one ``tests/test_fuzz.py`` checks with Hypothesis:
any ``ParsedDocument`` writes, and every text cell reads back as the text
the model held, with only the characters a worksheet cannot store replaced
by U+FFFD and carriage returns read as line feeds (XML 1.0 end-of-line
handling). Any exception, or a cell that reads back differently, is a crash
that libFuzzer reports and minimises.

Each input is decoded into a document: statements, entries and details,
with arbitrary text in the account owner, entry reference and
counterparty name, the fields that come straight from a bank.

Run (Linux x86_64, Python 3.12 to 3.14)::

    pip install --require-hashes -r requirements/fuzz.txt
    pip install --no-deps -e .
    python fuzz/fuzz_write_xlsx.py -max_total_time=60
"""

import re
import sys
import tempfile
from pathlib import Path
from typing import Any

import atheris

with atheris.instrument_imports():
    from camt053.models import (
        Account,
        Entry,
        ParsedDocument,
        Statement,
        TransactionDetails,
    )
    from openpyxl import load_workbook

    from camt053_writer_xlsx import write_xlsx

_UNWRITABLE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_OUT = Path(tempfile.mkdtemp(prefix="fuzz-xlsx-")) / "fuzz.xlsx"


def _read_back(value: str | None) -> str | None:
    """What a reader of the workbook sees for ``value``."""
    if not value:
        return None
    read = value.replace("\r\n", "\n").replace("\r", "\n")
    return _UNWRITABLE.sub("�", read)


def _text(fdp: "atheris.FuzzedDataProvider") -> str | None:
    return fdp.ConsumeUnicodeNoSurrogates(32) if fdp.ConsumeBool() else None


def _entry(fdp: "atheris.FuzzedDataProvider") -> Entry:
    details = [
        TransactionDetails(counterparty_name=_text(fdp))
        for _ in range(fdp.ConsumeIntInRange(0, 2))
    ]
    return Entry(reference=_text(fdp), amount=_text(fdp), details=details)


def _document(data: bytes) -> ParsedDocument:
    fdp = atheris.FuzzedDataProvider(data)
    statements = [
        Statement(
            id=_text(fdp),
            account=Account(owner_name=_text(fdp)),
            entries=[_entry(fdp) for _ in range(fdp.ConsumeIntInRange(0, 3))],
        )
        for _ in range(fdp.ConsumeIntInRange(0, 2))
    ]
    return ParsedDocument(message_type="camt.053.001.08", statements=statements)


def _column(workbook: Any, sheet: str, index: int) -> list[Any]:
    """Column ``index`` of ``sheet``, below the header."""
    rows = workbook[sheet].iter_rows(min_row=2, values_only=True)
    return [row[index] for row in rows]


def _counterparties(document: ParsedDocument) -> list[str | None]:
    """The Entries sheet's counterparty column, as a reader should see it."""
    return [
        _read_back(detail.counterparty_name if detail else None)
        for statement in document.statements
        for entry in statement.entries
        for detail in (entry.details or [None])
    ]


def test_one_input(data: bytes) -> None:
    """Write one document and hold the workbook to the contract."""
    document = _document(data)
    workbook = load_workbook(write_xlsx(document, _OUT))
    owners = [_read_back(s.account.owner_name) for s in document.statements]
    found = _column(workbook, "Metadata", 9)
    if found != owners:
        raise AssertionError(f"Metadata owners read back as {found!r}")
    found = _column(workbook, "Entries", 15)
    if found != _counterparties(document):
        raise AssertionError(f"Entries counterparties read back as {found!r}")


def main() -> None:
    """Hand the harness to libFuzzer."""
    atheris.Setup(sys.argv, test_one_input)
    atheris.Fuzz()


if __name__ == "__main__":
    main()
