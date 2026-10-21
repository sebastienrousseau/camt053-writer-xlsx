# SPDX-FileCopyrightText: 2026 Sebastien Rousseau <sebastian.rousseau@gmail.com>
# SPDX-License-Identifier: Apache-2.0 OR MIT
"""Property-based fuzzing of ``write_xlsx`` with Hypothesis.

The writer's input is a whole ``ParsedDocument``, and every text field in
it comes from a bank: camt.053 XML, or SWIFT MT940 and MT942 text, which
can carry characters XML cannot. For any document, the write must succeed
and every text cell must read back as the text the model held, with only
the characters a worksheet cannot store replaced by U+FFFD, and carriage
returns read as line feeds, as XML 1.0 end-of-line handling (section 2.11)
requires of every reader.
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

from camt053.models import (
    Account,
    Balance,
    Entry,
    ParsedDocument,
    Statement,
    TransactionDetails,
)
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from openpyxl import load_workbook

from camt053_writer_xlsx import write_xlsx

#: The characters openpyxl refuses, as the writer replaces them.
_UNWRITABLE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")

text = st.none() | st.text(max_size=40)


def _expected(value: str | None) -> str | None:
    """What a reader of the workbook sees for ``value``."""
    if value is None or value == "":
        return None
    read = value.replace("\r\n", "\n").replace("\r", "\n")
    return _UNWRITABLE.sub("\ufffd", read)


details = st.builds(
    TransactionDetails,
    end_to_end_id=text,
    counterparty_name=text,
    additional_info=text,
)
entries = st.builds(
    Entry,
    reference=text,
    amount=text,
    reason_code=text,
    reversal_indicator=st.booleans(),
    details=st.lists(details, max_size=2),
)
statements = st.builds(
    Statement,
    id=text,
    account=st.builds(Account, iban=text, owner_name=text),
    balances=st.lists(st.builds(Balance, type_code=text), max_size=2),
    entries=st.lists(entries, max_size=3),
)
documents = st.builds(
    ParsedDocument,
    message_type=st.just("camt.053.001.08"),
    msg_id=text,
    statements=st.lists(statements, max_size=2),
)


@settings(
    max_examples=150,
    deadline=None,
    suppress_health_check=[HealthCheck.too_slow],
)
@given(documents)
def test_any_document_writes_and_reads_back(document: ParsedDocument) -> None:
    """Every document writes, and its text cells hold the model's text."""
    with tempfile.TemporaryDirectory() as tmp:
        out = write_xlsx(document, Path(tmp) / "fuzz.xlsx")
        workbook = load_workbook(out)

    metadata = list(workbook["Metadata"].iter_rows(min_row=2, values_only=True))
    entry_rows = list(
        workbook["Entries"].iter_rows(min_row=2, values_only=True)
    )

    assert [row[1] for row in metadata] == [
        _expected(document.msg_id) for _ in document.statements
    ]
    assert [row[9] for row in metadata] == [
        _expected(s.account.owner_name) for s in document.statements
    ]
    expected_rows = sum(
        max(len(entry.details), 1)
        for statement in document.statements
        for entry in statement.entries
    )
    assert len(entry_rows) == expected_rows
    assert [row[1] for row in entry_rows] == [
        _expected(entry.reference)
        for statement in document.statements
        for entry in statement.entries
        for _ in range(max(len(entry.details), 1))
    ]
